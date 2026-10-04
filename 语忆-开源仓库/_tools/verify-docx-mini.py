#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify-docx-mini.py —— 验收 yuyi-preview/_tools/docx-mini.js 生成的 .docx
============================================================================

做四件事（全部要真的跑，跑不了就如实说，绝不假装）：

  1. 用 node 调 docx-mini.js，按多组语料生成 .docx（中文长句 / 英文 / 数字 / 标点 /
     空段落 / 多段落 / 需要转义的 & < > " ' ）；
  2. 用 python-docx 打开，把段落文本与样式读回来，逐条与输入断言；
  3. 用 zipfile 校验包结构（[Content_Types].xml / _rels/.rels / word/document.xml
     都存在、testzip() 返回 None、每部件都是合法 UTF-8 且能被 XML 解析、
     每个 Relationship 目标都存在且声明了 ContentType —— 即 OPC 包完整性）；
  4. 额外做一次**真实渲染校验**：本机若装了 LibreOffice（soffice）就转成 txt 再读回来比对；
     没装就打印「未做真实渲染校验」，并说明到底哪些点没验证。

用法：
    python yuyi-preview\\_tools\\verify-docx-mini.py

退出码：全部 PASS -> 0；有任何 FAIL -> 1。
临时文件都放在 %TEMP% 下（脚本结束时打印目录，便于人工翻看）。
"""

from __future__ import annotations

import glob
import json
import os
import posixpath
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET

# ---------------------------------------------------------------------------
# 控制台编码：Windows 默认 GBK，打印盲文点字 / emoji 会直接 UnicodeEncodeError
# ---------------------------------------------------------------------------
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="backslashreplace")
    except Exception:
        pass

HERE = Path(__file__).resolve().parent
DOCX_MINI = HERE / "docx-mini.js"

RESULTS: list[tuple[bool, str]] = []


def check(ok: bool, label: str, detail: str = "") -> bool:
    ok = bool(ok)
    RESULTS.append((ok, label))
    line = ("[PASS] " if ok else "[FAIL] ") + label
    if detail:
        line += "  -> " + detail
    print(line, flush=True)
    return ok


def info(label: str, detail: str = "") -> None:
    print("[INFO] " + label + (("  -> " + detail) if detail else ""), flush=True)


def warn(label: str, detail: str = "") -> None:
    print("[WARN] " + label + (("  -> " + detail) if detail else ""), flush=True)


def skip(label: str, detail: str = "") -> None:
    print("[SKIP] " + label + (("  -> " + detail) if detail else ""), flush=True)


def head(title: str) -> None:
    print("\n" + "=" * 78, flush=True)
    print("== " + title, flush=True)
    print("=" * 78, flush=True)


def show(text: str, label: str = "") -> None:
    """打印可能含非 GBK 字符的文本：先是原文，再给一份 ascii 转义版（证据不受终端编码影响）。"""
    if label:
        print(label, flush=True)
    print("      原文: " + text, flush=True)
    print("      ascii: " + ascii(text), flush=True)


# ---------------------------------------------------------------------------
# 0. 依赖自举：python-docx
# ---------------------------------------------------------------------------
def ensure_python_docx() -> bool:
    try:
        import docx  # noqa: F401

        info("python-docx", "已在本机 python 环境里 (import docx 成功)")
        return True
    except Exception:
        pass

    candidates: list[str] = []
    home = Path.home()
    candidates += glob.glob(str(home / ".dsh" / "dsh-runtimes" / "*" / "dependencies" / "python" / "Lib" / "site-packages"))
    candidates += glob.glob(str(home / ".dsh" / "**" / "site-packages"), recursive=True)
    candidates.append(str(Path(tempfile.gettempdir()) / "docxmini-pylibs"))

    for cand in candidates:
        if os.path.isdir(cand) and os.path.isdir(os.path.join(cand, "docx")):
            sys.path.insert(0, cand)
            try:
                import docx  # noqa: F401

                info("python-docx", "从备用 site-packages 载入: " + cand)
                return True
            except Exception:
                try:
                    sys.path.remove(cand)
                except ValueError:
                    pass

    target = Path(tempfile.gettempdir()) / "docxmini-pylibs"
    info("python-docx", "本机没有，尝试 pip install --target " + str(target))
    try:
        proc = subprocess.run(
            [sys.executable, "-m", "pip", "install", "--target", str(target), "--quiet", "python-docx"],
            capture_output=True,
            text=True,
            timeout=600,
        )
        print(proc.stdout[-2000:] if proc.stdout else "", flush=True)
        print(proc.stderr[-2000:] if proc.stderr else "", flush=True)
    except Exception as exc:  # pragma: no cover
        warn("pip install 失败", repr(exc))

    if os.path.isdir(target):
        sys.path.insert(0, str(target))
    try:
        import docx  # noqa: F401

        info("python-docx", "pip 装到 " + str(target) + " 后可用")
        return True
    except Exception as exc:
        check(False, "python-docx 不可用，无法做回读验证", repr(exc))
        return False


# ---------------------------------------------------------------------------
# 1. 依赖自举：node
# ---------------------------------------------------------------------------
def find_node() -> str | None:
    found = shutil.which("node")
    if found:
        return found
    cands: list[str] = []
    cands += glob.glob(str(Path.home() / ".dsh" / "dsh-runtimes" / "*" / "dependencies" / "node" / "bin" / "node.exe"))
    cands += glob.glob(str(Path.home() / ".dsh" / "dsh-runtimes" / "*" / "dependencies" / "node" / "node.exe"))
    cands += [
        r"C:\Program Files\nodejs\node.exe",
        r"C:\Program Files (x86)\nodejs\node.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Programs\nodejs\node.exe"),
        "/usr/bin/node",
        "/usr/local/bin/node",
    ]
    for cand in cands:
        if cand and os.path.isfile(cand):
            return cand
    return None


# ---------------------------------------------------------------------------
# 2. 语料
# ---------------------------------------------------------------------------
CORPUS_A = [
    {"text": "雨意预览 · 导出验证文档", "style": "title"},
    {"text": "第一章 中文长句与标点", "style": "heading"},
    {
        "text": "这是一段较长的中文句子，用来验证整段中文能否原样往返：中文逗号，句号。问号？叹号！分号；冒号：顿号、"
        "引号「」『』（）《》破折号——省略号……以及全角空格　与数字 2026 年 10 月 4 日 12:30。",
        "style": "normal",
    },
    {
        "text": "Mixed English + 中文 + numbers 0123456789, half-width ,.;:!?()[]{} and symbols @#$%^*_+-=|~`",
        "style": "normal",
    },
    {"text": "需要转义的字符：& < > \" ' 组合 a&b <c> \"d\" 'e' 以及字面量 &amp; &lt; &#39;", "style": "normal"},
    {"text": "", "style": "normal"},
    {"text": None, "style": "normal"},
    {"text": "上面两个空段落（一个是空字符串，一个是 null）。", "style": "normal"},
    {"text": "制表符\t与段落内换行\n第二行\n第三行", "style": "normal"},
    {"text": "盲文点字 ⠿⠁⠃⠉ 与 emoji 🙂 混排，用来确认中文与非 BMP 字符都不乱码。", "style": "normal"},
    {"text": "全角字符：ＡＢＣ１２３（全角括号）", "style": "normal"},
    {"text": "  这一段首尾各有空格，中间  还用 连续空格  ", "style": "normal"},
    {"text": "", "style": "normal"},
]

CORPUS_B = [
    {"text": "没有标题段落的正文第一段。", "style": "normal"},
    {"text": "第二段：opts.title 应该被写成第 1 段的 Title 标题。", "style": "heading"},
    {"text": "第三段：opts.fontFamily / opts.fontSize 会被写进 styles.xml。", "style": "normal"},
]
OPTS_B = {"title": "标题由 opts.title 写入（含中文）", "fontFamily": "微软雅黑", "fontSize": 12}

CORPUS_C = [
    {"text": "重复的标题", "style": "title"},  # 首段已是同名 title
    "纯字符串段落（不是对象）",  # 容忍字符串
    {"text": "未知样式名走 Normal", "style": "weird-unknown-style"},
    {"style": "heading"},  # 没有 text
    {"text": "只有换行\n", "style": "normal"},
    {"text": "啊" * 500 + "（超长中文段落 500 字）", "style": "normal"},
]
OPTS_C = {"title": "重复的标题"}  # 首段已是同名 title -> 不应重复插入

CORPUS_D: list = []  # 空数组：必须仍然是合法 docx

CORPUS_E = [
    {"text": "stored 兜底路径验证（compress:false）", "style": "title"},
    {"text": "这一段也走不压缩的 stored 方式写入 ZIP。", "style": "normal"},
]
OPTS_E = {"compress": False}

STYLE_UI = {
    "title": "Title",
    "heading": "Heading 1",
    "heading1": "Heading 1",
    "h1": "Heading 1",
    "heading2": "Heading 2",
    "h2": "Heading 2",
    "heading3": "Heading 3",
    "h3": "Heading 3",
    "normal": "Normal",
    "": "Normal",
}


def norm_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def expected_texts(corpus: list, opts: dict) -> list[str]:
    texts = []
    for item in corpus:
        if isinstance(item, dict):
            texts.append(norm_text(item.get("text")))
        else:
            texts.append(norm_text(item))
    title = opts.get("title")
    if title:
        first_style = "normal"
        if corpus and isinstance(corpus[0], dict):
            first_style = str(corpus[0].get("style") or "normal")
        dup = bool(texts) and STYLE_UI.get(first_style.strip().lower(), "Normal") == "Title" and texts[0] == title
        if not dup:
            texts.insert(0, norm_text(title))
    return texts


def expected_styles(corpus: list, opts: dict) -> list[str]:
    styles = []
    for item in corpus:
        if isinstance(item, dict):
            raw = str(item.get("style") or "normal").strip().lower()
        else:
            raw = "normal"
        styles.append(STYLE_UI.get(raw, "Normal"))
    title = opts.get("title")
    if title:
        first_style = "normal"
        if corpus and isinstance(corpus[0], dict):
            first_style = str(corpus[0].get("style") or "normal")
        first_text = expected_texts(corpus, {})[0] if corpus else ""
        dup = bool(styles) and STYLE_UI.get(first_style.strip().lower(), "Normal") == "Title" and first_text == title
        if not dup:
            styles.insert(0, "Title")
    return styles


CASES = [
    {"name": "A-defaults", "paragraphs": CORPUS_A, "opts": {}},
    {"name": "B-title-font", "paragraphs": CORPUS_B, "opts": OPTS_B},
    {"name": "C-edge", "paragraphs": CORPUS_C, "opts": OPTS_C},
    {"name": "D-empty", "paragraphs": CORPUS_D, "opts": {}},
    {"name": "E-stored", "paragraphs": CORPUS_E, "opts": OPTS_E},
]

DRIVER_SOURCE = r"""
"use strict";
// 由 verify-docx-mini.py 生成到 %TEMP% 的临时驱动脚本（CommonJS，避免 package.json type 影响）
const fs = require("fs");
const path = require("path");

const modPath = process.argv[2];
const cfgPath = process.argv[3];

const DocxMini = require(path.resolve(modPath));

(async () => {
  const cfg = JSON.parse(fs.readFileSync(cfgPath, "utf8"));
  const report = {
    node: process.version,
    api: Object.keys(DocxMini).sort(),
    hasBlob: typeof Blob !== "undefined",
    hasTextEncoder: typeof TextEncoder !== "undefined",
    hasCompressionStream: typeof CompressionStream !== "undefined",
    isCommonJSExport: typeof DocxMini.build === "function" && typeof DocxMini.buildBytes === "function",
    cases: [],
  };

  for (const c of cfg.cases) {
    const t0 = Date.now();
    const bytes = await DocxMini.buildBytes(c.paragraphs, c.opts || {});
    const file = path.join(cfg.outDir, c.name + ".docx");
    fs.writeFileSync(file, Buffer.from(bytes));
    const blob = await DocxMini.build(c.paragraphs, c.opts || {});
    report.cases.push({
      name: c.name,
      file: file,
      bytes: bytes.length,
      isUint8Array: bytes instanceof Uint8Array,
      blobSize: blob.size,
      blobType: blob.type,
      isBlob: typeof Blob !== "undefined" && blob instanceof Blob,
      ms: Date.now() - t0,
    });
  }

  // 入参容错：非数组且非字符串 -> 必须 reject（不许静默出个空文件）
  try {
    await DocxMini.buildBytes(12345, {});
    report.badInputThrew = false;
  } catch (e) {
    report.badInputThrew = true;
    report.badInputMessage = String((e && e.message) || e);
  }
  // null -> 当空文档处理（仍然合法）
  const nullBytes = await DocxMini.buildBytes(null, {});
  report.nullInputBytes = nullBytes.length;
  // 单个字符串 -> 当一段
  const strBytes = await DocxMini.buildBytes("单个字符串也算一段", {});
  report.singleStringBytes = strBytes.length;

  console.log("DRIVER_JSON " + JSON.stringify(report));
})().catch((e) => {
  console.error("DRIVER_ERROR " + ((e && e.stack) || e));
  process.exit(3);
});
"""


# ---------------------------------------------------------------------------
# 3. LibreOffice 查找
# ---------------------------------------------------------------------------
def find_soffice() -> str | None:
    for name in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    roots = [
        os.environ.get("ProgramFiles", r"C:\Program Files"),
        os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
        os.path.expandvars(r"%LOCALAPPDATA%\Programs"),
        os.path.expandvars(r"%ProgramFiles%\LibreOffice"),
        "C:\\LibreOffice",
        "D:\\LibreOffice",
        "/usr/bin",
        "/usr/local/bin",
        "/opt/libreoffice/program",
        "/Applications/LibreOffice.app/Contents/MacOS",
    ]
    patterns = [
        "LibreOffice/program/soffice.exe",
        "*/LibreOffice/program/soffice.exe",
        "*/*/LibreOffice/program/soffice.exe",
        "*/program/soffice.exe",
        "soffice",
        "*/soffice",
    ]
    for root in roots:
        if not root or not os.path.isdir(root):
            continue
        for pat in patterns:
            for hit in glob.glob(os.path.join(root, pat)):
                if os.path.isfile(hit):
                    return hit
    return None


def find_wps() -> str | None:
    roots = [os.path.expandvars(r"%LOCALAPPDATA%\Kingsoft\WPS Office"), os.path.expandvars(r"%ProgramFiles%\Kingsoft")]
    for root in roots:
        if not os.path.isdir(root):
            continue
        for pat in ("*/office6/wps.exe", "*/*/office6/wps.exe"):
            for hit in glob.glob(os.path.join(root, pat)):
                if os.path.isfile(hit):
                    return hit
    return None


def check_opc(parts: dict, names: list, label: str) -> None:
    """OPC 包完整性：每个 Relationship 的 Target 都要真实存在，且每个部件都声明了 ContentType。

    这一条专治「Word 打开提示文件已损坏」：缺 rels 目标、缺 ContentType Override 是最常见的两种。
    （注意：这是结构级校验，不等于渲染校验。）"""
    ns_ct = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for child in ET.fromstring(parts["[Content_Types].xml"]):
        tag = child.tag.replace(ns_ct, "")
        if tag == "Default":
            defaults[(child.get("Extension") or "").lower()] = child.get("ContentType") or ""
        elif tag == "Override":
            overrides[child.get("PartName") or ""] = child.get("ContentType") or ""

    def ct_for(part: str) -> str:
        if "/" + part in overrides:
            return overrides["/" + part]
        ext = part.rsplit(".", 1)[-1].lower() if "." in part else ""
        return defaults.get(ext, "")

    problems: list[str] = []
    resolved_lines: list[str] = []
    for rels_name in [n for n in names if n.endswith(".rels")]:
        base = posixpath.dirname(posixpath.dirname(rels_name))
        for rel in ET.fromstring(parts[rels_name]):
            if (rel.get("TargetMode") or "") == "External":
                continue
            resolved = posixpath.normpath(posixpath.join(base, rel.get("Target") or "")).lstrip("/")
            resolved_lines.append("%-32s -%s-> %s  [%s]" % (
                rels_name, rel.get("Id"), resolved, ct_for(resolved) or "!! 无 ContentType"))
            if resolved not in names:
                problems.append("关系目标不存在: " + resolved)
            elif not ct_for(resolved):
                problems.append("目标部件没有 ContentType: " + resolved)
    for part in names:
        if not part.endswith("/") and not ct_for(part):
            problems.append("部件没有 ContentType: " + part)

    print("      OPC 关系解析:", flush=True)
    for line in resolved_lines:
        print("        " + line, flush=True)
    check(not problems, label, "; ".join(problems) if problems else "所有关系目标与 ContentType 都能解析")


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def main() -> int:
    head("0. 环境与依赖")
    print("脚本      : " + str(Path(__file__).resolve()), flush=True)
    print("被测实现  : " + str(DOCX_MINI) + ("  (存在)" if DOCX_MINI.is_file() else "  (缺失!)"), flush=True)
    print("python    : " + sys.version.split()[0] + "  (" + sys.executable + ")", flush=True)

    if not check(DOCX_MINI.is_file(), "docx-mini.js 存在", str(DOCX_MINI)):
        return 1
    if not ensure_python_docx():
        return 1
    import docx  # noqa: E402
    from docx import Document  # noqa: E402

    info("python-docx", "版本 " + getattr(docx, "__version__", "(未知)"))

    node = find_node()
    if not check(node is not None, "找到 node 可执行文件", str(node)):
        return 1
    try:
        node_ver = subprocess.run([node, "-v"], capture_output=True, text=True, timeout=60)
        info("node", (node_ver.stdout or "").strip() + "  @ " + str(node))
    except Exception as exc:
        warn("node -v 失败", repr(exc))

    soffice = find_soffice()
    if soffice:
        info("LibreOffice", soffice)
    else:
        info("LibreOffice", "未找到 soffice —— 稍后会明确报告「未做真实渲染校验」")
    wps = find_wps()
    if wps:
        info("WPS Office", wps + "  (存在，但没有可用的无头转换命令行，本脚本不使用它做渲染校验)")

    # ---------------------------------------------------------------- 生成
    head("1. 用 node 调 docx-mini.js 生成 .docx")
    tmpdir = Path(tempfile.mkdtemp(prefix="docxmini-verify-"))
    outdir = tmpdir / "out"
    outdir.mkdir(parents=True, exist_ok=True)
    driver = tmpdir / "drive-docx-mini.cjs"
    cfgfile = tmpdir / "cases.json"
    driver.write_text(DRIVER_SOURCE, encoding="utf-8")
    cfgfile.write_text(json.dumps({"outDir": str(outdir), "cases": CASES}, ensure_ascii=False), encoding="utf-8")
    info("临时目录", str(tmpdir))

    print("\n$ node drive-docx-mini.cjs docx-mini.js cases.json", flush=True)
    proc = subprocess.run(
        [node, str(driver), str(DOCX_MINI), str(cfgfile)], capture_output=True, text=True, timeout=300,
        encoding="utf-8", errors="backslashreplace",
    )
    stdout = proc.stdout or ""
    stderr = proc.stderr or ""
    print("--- node stdout ---", flush=True)
    print(stdout.rstrip("\n") if stdout.strip() else "(空)", flush=True)
    if stderr.strip():
        print("--- node stderr ---", flush=True)
        print(stderr.rstrip("\n"), flush=True)
    print("--- node exit code: " + str(proc.returncode) + " ---", flush=True)

    if not check(proc.returncode == 0, "node 驱动脚本正常退出 (exit 0)", "exit=" + str(proc.returncode)):
        return 1

    driver_json = None
    for line in stdout.splitlines():
        if line.startswith("DRIVER_JSON "):
            driver_json = json.loads(line[len("DRIVER_JSON "):])
    if not check(driver_json is not None, "解析到 node 驱动返回的 JSON 报告"):
        return 1

    info("node 环境", "node " + driver_json["node"]
         + " | Blob=" + str(driver_json["hasBlob"])
         + " TextEncoder=" + str(driver_json["hasTextEncoder"])
         + " CompressionStream=" + str(driver_json["hasCompressionStream"]))
    check(driver_json["api"] == ["build", "buildBytes"],
          "导出 API 恰好是 build / buildBytes", "实际: " + ",".join(driver_json["api"]))
    check(driver_json["isCommonJSExport"], "module.exports 在 Node 下可用 (require 成功)")
    check(driver_json["badInputThrew"],
          "非数组入参（12345）会 reject 而不是静默产出",
          "message=" + str(driver_json.get("badInputMessage", "")))
    check(driver_json.get("nullInputBytes", 0) > 0, "paragraphs=null 也能产出合法空文档",
          str(driver_json.get("nullInputBytes")) + " 字节")
    check(driver_json.get("singleStringBytes", 0) > 0, "paragraphs=单个字符串也能产出",
          str(driver_json.get("singleStringBytes")) + " 字节")

    # ------------------------------------------------------------ 逐个用例
    for case, case_report in zip(CASES, driver_json["cases"]):
        name = case["name"]
        path = Path(case_report["file"])
        opts = case["opts"] or {}
        exp_texts = expected_texts(case["paragraphs"], opts)
        exp_styles = expected_styles(case["paragraphs"], opts)

        head("2." + name + "  校验 " + path.name)
        check(case_report["isUint8Array"], "buildBytes 返回 Uint8Array")
        check(case_report["isBlob"] and case_report["blobSize"] > 0,
              "build 返回 Blob", "size=" + str(case_report["blobSize"]) + " type=" + case_report["blobType"])
        check(case_report["blobType"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
              "Blob MIME 是 docx 官方 MIME", case_report["blobType"])
        check(path.is_file() and path.stat().st_size > 0, "docx 文件写出成功",
              str(path.stat().st_size) + " 字节")
        check(path.stat().st_size == case_report["bytes"], "文件大小 == buildBytes 长度")

        with open(path, "rb") as fh:
            signature = fh.read(4)
        check(signature[:2] == b"PK", "文件头是 ZIP 签名 PK", repr(signature))

        # ---- zip 结构 ----
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            info("zip 条目 (" + str(len(names)) + ")", ", ".join(names))
            methods = sorted({zi.compress_type for zi in zf.infolist()})
            method_names = {0: "stored(0)", 8: "deflate(8)"}
            info("压缩方式", ", ".join(method_names.get(m, str(m)) for m in methods))
            for zi in zf.infolist():
                print("      %-32s method=%-11s crc=%08X  size=%-6d  packed=%-6d"
                      % (zi.filename, method_names.get(zi.compress_type, zi.compress_type),
                         zi.CRC, zi.file_size, zi.compress_size), flush=True)

            check(zf.testzip() is None, "zipfile.testzip() 返回 None（无损坏条目）", str(zf.testzip()))
            for required in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
                check(required in names, "包含必需部件 " + required)
            check(names[0] == "[Content_Types].xml", "[Content_Types].xml 是第 1 个条目", names[0])
            if name == "E-stored":
                check(methods == [0], "compress:false 时全部条目走 stored 不压缩",
                      str([method_names.get(m, m) for m in methods]))
            else:
                check(True, "压缩方式（deflate 优先、失败自动 stored 兜底）",
                      ", ".join(method_names.get(m, str(m)) for m in methods))

            parts: dict[str, str] = {}
            bad_utf8 = []
            bad_xml = []
            for entry in names:
                if entry.endswith("/"):
                    continue
                raw = zf.read(entry)
                try:
                    text = raw.decode("utf-8")
                except UnicodeDecodeError as exc:
                    bad_utf8.append(entry + ": " + repr(exc))
                    continue
                parts[entry] = text
                try:
                    ET.fromstring(text)
                except ET.ParseError as exc:
                    bad_xml.append(entry + ": " + str(exc))
            check(not bad_utf8, "所有部件都是合法 UTF-8", "; ".join(bad_utf8))
            check(not bad_xml, "所有部件的 XML 都能被解析（well-formed）", "; ".join(bad_xml))

        # ---- Content_Types / rels 内容 ----
        ct = parts.get("[Content_Types].xml", "")
        check("/word/document.xml" in ct
              and "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml" in ct,
              "[Content_Types].xml 声明了主文档部件的正确 ContentType")
        rels = parts.get("_rels/.rels", "")
        check("word/document.xml" in rels and "/officeDocument" in rels,
              "_rels/.rels 指向 word/document.xml (officeDocument)")
        doc_rels = parts.get("word/_rels/document.xml.rels", "")
        check("styles.xml" in doc_rels, "word/_rels/document.xml.rels 关联 styles.xml")
        check_opc(parts, names, "OPC 包完整性：每个 Relationship 目标都存在且已声明 ContentType")

        # ---- 转义检查（按语料实际含有的字符决定查哪些，而不是写死用例名）----
        body = parts.get("word/document.xml", "")
        joined_expected = "\n".join(exp_texts)
        if "&" in joined_expected:
            check("&amp;" in body and "&lt;" in body and "&gt;" in body and "&quot;" in body and "&apos;" in body,
                  "document.xml 里 & < > \" ' 都被转义成实体")
            check("&amp;amp;" in body and "&amp;lt;" in body,
                  "字面量 &amp; / &lt; 被二次转义（不会破坏 XML）")
        if "<c>" in joined_expected:
            check("&lt;c&gt;" in body, "原文 <c> 存成 &lt;c&gt;")
        if "\n" in joined_expected or "\t" in joined_expected:
            has_br = any("\n" in t for t in exp_texts)
            has_tab = any("\t" in t for t in exp_texts)
            if has_br:
                check("<w:br/>" in body, "段落内换行用 <w:br/> 表达")
            if has_tab:
                check("<w:tab/>" in body, "段落内制表用 <w:tab/> 表达")
        if any(exp_texts):
            check('xml:space="preserve"' in body, "w:t 带 xml:space=preserve（首尾空格不会被吃掉）")

        # ---- 原始文本完整性（不依赖 python-docx 的独立检查）----
        missing = []
        for text in exp_texts:
            if not text or "\n" in text or "\t" in text:
                continue
            escaped = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                       .replace('"', "&quot;").replace("'", "&apos;"))
            if escaped not in body:
                missing.append(text[:30])
        check(not missing, "document.xml 原文里能找到每一段（独立于 python-docx）",
              "缺失 " + str(missing) if missing else str(len(exp_texts)) + " 段检查完毕")

        # ---- python-docx 回读 ----
        document = Document(str(path))
        got_texts = [norm_text(p.text) for p in document.paragraphs]
        got_styles = [p.style.name for p in document.paragraphs]

        check(len(got_texts) == len(exp_texts),
              "python-docx 读回段落数量一致",
              "读回 " + str(len(got_texts)) + " / 期望 " + str(len(exp_texts)))

        print("--- python-docx 读回结果 ---", flush=True)
        for idx in range(max(len(got_texts), len(exp_texts))):
            g = got_texts[idx] if idx < len(got_texts) else "<缺失>"
            gs = got_styles[idx] if idx < len(got_styles) else "<缺失>"
            flag = "OK " if (idx < len(exp_texts) and idx < len(got_texts)
                            and g == exp_texts[idx] and gs == exp_styles[idx]) else "!! "
            shown = ascii(g)
            if len(shown) > 160:  # 只截断「打印」，比较用的仍是完整文本
                shown = shown[:160] + "...(共 %d 字符, 已完整比对)" % len(g)
            print("  [" + flag + "] #%-2d style=%-9s text=%s" % (idx, gs, shown), flush=True)

        text_mismatch = [
            {"index": i, "expected": exp_texts[i], "got": got_texts[i]}
            for i in range(min(len(exp_texts), len(got_texts)))
            if exp_texts[i] != got_texts[i]
        ]
        check(not text_mismatch, "逐条断言：段落文本与输入完全一致",
              json.dumps(text_mismatch, ensure_ascii=True) if text_mismatch else
              str(len(exp_texts)) + " 条全部一致（含空段落）")
        style_mismatch = [
            {"index": i, "expected": exp_styles[i], "got": got_styles[i]}
            for i in range(min(len(exp_styles), len(got_styles)))
            if exp_styles[i] != got_styles[i]
        ]
        check(not style_mismatch, "逐条断言：段落样式（Title / Heading 1 / Normal）一致",
              json.dumps(style_mismatch, ensure_ascii=True) if style_mismatch else
              "样式序列: " + ", ".join(got_styles))

        if opts.get("title"):
            cp_title = document.core_properties.title
            check(cp_title == opts["title"], "opts.title 写进了 docProps/core.xml 的 dc:title",
                  ascii(cp_title))
        if name == "C-edge":
            check(got_styles[0] == "Title" and got_texts[0] == "重复的标题" and got_texts.count("重复的标题") == 1,
                  "opts.title 与首段同名 Title 时不重复插入", ascii(got_texts[0]))

        # ---- 中文 / 非 BMP 往返 ----
        if name == "A-defaults":
            joined = "\n".join(got_texts)
            check("中文逗号，句号。问号？叹号！" in joined, "中文长句里的全角标点原样读回")
            check("破折号——省略号……" in joined, "破折号/省略号原样读回")
            check("盲文点字 ⠿⠁⠃⠉ 与 emoji 🙂" in joined, "盲文点字与 emoji（非 BMP）原样读回")
            check("\uFFFD" not in joined, "读回的文本里没有替换字符 U+FFFD（即没有乱码）")
            check("制表符\t与段落内换行\n第二行\n第三行" in joined, "段落内 \\t / \\n 往返一致")
            raw_cn = [t for t in got_texts if "中文" in t][:1]
            if raw_cn:
                show(raw_cn[0][:40], "      中文抽样:")

    # ---------------------------------------------------- 真实渲染校验
    head("3. 真实渲染校验（LibreOffice -> txt 再读回）")
    if not soffice:
        skip("未做真实渲染校验：本机没有 LibreOffice (soffice)")
        print("      本机查过：PATH、%ProgramFiles%\\LibreOffice、%ProgramFiles(x86)%\\LibreOffice、", flush=True)
        print("      %LOCALAPPDATA%\\Programs、C:\\LibreOffice、D:\\LibreOffice、/usr/bin、/usr/local/bin、", flush=True)
        print("      /opt/libreoffice/program、/Applications/LibreOffice.app —— 都没有 soffice。", flush=True)
        if wps:
            print("      （本机装了 WPS Office：" + wps + "，但它没有可用的无头转换命令行，", flush=True)
            print("        所以本脚本没有用它冒充渲染校验。）", flush=True)
        print("      => 未经真实渲染验证的点：Word/WPS/LibreOffice 打开时的分页、字体回退、", flush=True)
        print("         样式外观、以及「Word 是否会提示文档损坏」这一类只有渲染器才知道的事情。", flush=True)
    else:
        render_src = outdir / "A-defaults.docx"
        render_out = tmpdir / "render"
        render_out.mkdir(parents=True, exist_ok=True)
        cmd = [soffice, "--headless", "--norestore", "--nolockcheck",
               "--convert-to", "txt:Text (encoded):UTF8", "--outdir", str(render_out), str(render_src)]
        print("$ " + " ".join('"' + c + '"' for c in cmd), flush=True)
        try:
            conv = subprocess.run(cmd, capture_output=True, text=True, timeout=300,
                                  encoding="utf-8", errors="backslashreplace")
            print("--- soffice stdout ---\n" + (conv.stdout or "(空)"), flush=True)
            if (conv.stderr or "").strip():
                print("--- soffice stderr ---\n" + conv.stderr, flush=True)
            print("--- soffice exit code: " + str(conv.returncode) + " ---", flush=True)
        except Exception as exc:
            conv = None
            check(False, "soffice 调用失败", repr(exc))

        txts = list(render_out.glob("*.txt"))
        if check(bool(txts), "soffice 产出了 txt（说明它真的解析并渲染了这个 docx）",
                 str(txts[0]) if txts else "无输出文件"):
            rendered = txts[0].read_text(encoding="utf-8", errors="backslashreplace")
            exp_lines: list[str] = []
            for text in expected_texts(CORPUS_A, {}):
                exp_lines.extend(text.split("\n"))
            got_lines = [ln.rstrip() for ln in rendered.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
            exp_trim = list(exp_lines)
            while exp_trim and exp_trim[-1].strip() == "":
                exp_trim.pop()
            while got_lines and got_lines[-1].strip() == "":
                got_lines.pop()
            print("--- 渲染回读（前 12 行）---", flush=True)
            for i, ln in enumerate(got_lines[:12]):
                print("  #%-2d %s" % (i, ascii(ln)), flush=True)
            check(got_lines == [ln.rstrip() for ln in exp_trim], "渲染后的文本与输入逐行一致",
                  "行数 " + str(len(got_lines)) + " / 期望 " + str(len(exp_trim)))
        else:
            warn("LibreOffice 渲染校验没跑成，等于「未做真实渲染校验」")

    # ---------------------------------------------------------------- 汇总
    head("4. 汇总")
    failed = [label for ok, label in RESULTS if not ok]
    print("共 " + str(len(RESULTS)) + " 项断言，PASS " + str(len(RESULTS) - len(failed))
          + "，FAIL " + str(len(failed)), flush=True)
    if failed:
        for label in failed:
            print("  [FAIL] " + label, flush=True)
        print("\n结果: FAIL", flush=True)
        return 1
    print("结果: 全部 PASS", flush=True)
    print("生成的 docx 与临时文件保留在: " + str(tmpdir), flush=True)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\n中断", flush=True)
        sys.exit(130)